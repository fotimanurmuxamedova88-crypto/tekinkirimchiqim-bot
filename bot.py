import asyncio
import logging
import os
import re
from contextlib import closing
from io import BytesIO

import pymysql
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from telegram.ext import ApplicationBuilder, CommandHandler, MessageHandler, filters

LOG = logging.getLogger("finance_bot")
HELP = (
    "✅ Kirim-Chiqim bot\n"
    "Kirim: savdo 500000\n"
    "Chiqim: gaz -120000\n"
    "Summa butun so'mda, bo'sh joysiz yoziladi.\n"
    "/balans — balans\n/hisobot — kirim va chiqim\n"
    "/excel yoki /exel — Excel fayl"
)


def load_config():
    required = ("BOT_TOKEN", "ALLOWED_USER_ID", "MYSQLHOST", "MYSQLUSER",
                "MYSQLPASSWORD", "MYSQLDATABASE")
    missing = [key for key in required if not os.environ.get(key)]
    if missing:
        raise ValueError("Missing environment variables: " + ", ".join(missing))
    try:
        owner = int(os.environ["ALLOWED_USER_ID"])
        port = int(os.environ.get("MYSQLPORT", "3306"))
    except ValueError:
        raise ValueError("ALLOWED_USER_ID and MYSQLPORT must be integers") from None
    if owner <= 0 or not 1 <= port <= 65535:
        raise ValueError("Invalid ALLOWED_USER_ID or MYSQLPORT range")
    return {
        "token": os.environ["BOT_TOKEN"].strip(), "owner": owner,
        "db": dict(host=os.environ["MYSQLHOST"], user=os.environ["MYSQLUSER"],
                   password=os.environ["MYSQLPASSWORD"],
                   database=os.environ["MYSQLDATABASE"], port=port,
                   charset="utf8mb4", connect_timeout=10,
                   read_timeout=15, write_timeout=15, autocommit=False),
    }


def query(db, sql, params=(), *, write=False):
    with closing(pymysql.connect(**db)) as conn:
        try:
            with conn.cursor() as cur:
                cur.execute(sql, params)
                result = None if write else cur.fetchall()
            if write:
                conn.commit()
            return result
        except Exception:
            conn.rollback()
            raise


async def initialize(app):
    db = app.bot_data["config"]["db"]
    for attempt in range(5):
        try:
            await asyncio.to_thread(query, db, """
                CREATE TABLE IF NOT EXISTS transactions (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    nomi VARCHAR(255),
                    summa BIGINT
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """, write=True)
            break
        except pymysql.OperationalError as exc:
            code = exc.args[0] if exc.args else None
            LOG.error("Database initialization failed: code=%s", code)
            if code not in (2002, 2003, 2006, 2013) or attempt == 4:
                raise
            await asyncio.sleep(min(2 ** attempt, 8))
    LOG.info("Database initialized; starting polling")


def authorized(update, context):
    return bool(update.effective_user and update.effective_chat
                and update.effective_chat.type == "private"
                and update.effective_user.id == context.bot_data["config"]["owner"])


def parse_operation(text):
    try:
        name, raw = text.strip().rsplit(maxsplit=1)
    except ValueError:
        raise ValueError("Misol: savdo 500000 yoki gaz -120000") from None
    if not re.fullmatch(r"[+-]?(?:[0-9]+|[0-9]{1,3}(?:,[0-9]{3})+)", raw):
        raise ValueError("Summa butun son bo'lsin: 120000 yoki -120000")
    amount = int(raw.replace(",", ""))
    if amount == 0 or not -(2 ** 63) <= amount < 2 ** 63:
        raise ValueError("Summa nol yoki ruxsat etilgan chegaradan tashqarida")
    if not name or len(name) > 255 or ILLEGAL_CHARACTERS_RE.search(name):
        raise ValueError("Nomi 1–255 ta belgi bo'lsin; boshqaruv belgilarisiz")
    return name, amount


def make_excel(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kirim-Chiqim"
    ws.append(["Nomi", "Summa"])
    for name, amount in rows:
        name = ILLEGAL_CHARACTERS_RE.sub("", str(name or ""))
        # Excel numeric cells preserve only 15 significant decimal digits.
        value = str(amount) if amount is not None and abs(amount) >= 10 ** 15 else amount
        ws.append([name, value])
        # Product names are text, even when they begin with '='.
        ws.cell(ws.max_row, 1).data_type = "s"
        ws.cell(ws.max_row, 2).number_format = '#,##0;[Red]-#,##0'
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 22
    with BytesIO() as stream:
        wb.save(stream)
        result = stream.getvalue()
    wb.close()
    return result


async def start(update, context):
    if authorized(update, context):
        await update.effective_message.reply_text(HELP)


async def report(update, context):
    if not authorized(update, context):
        return
    rows = await asyncio.to_thread(query, context.bot_data["config"]["db"], """
        SELECT COALESCE(SUM(CASE WHEN summa > 0 THEN summa ELSE 0 END), 0),
               COALESCE(SUM(CASE WHEN summa < 0 THEN -CAST(summa AS DECIMAL(30,0)) ELSE 0 END), 0)
        FROM transactions
    """)
    income, expense = rows[0]
    await update.effective_message.reply_text(
        f"Kirim: {income:,} so'm\nChiqim: {expense:,} so'm\n"
        f"Balans: {income - expense:,} so'm")


async def excel(update, context):
    if not authorized(update, context):
        return
    rows = await asyncio.to_thread(query, context.bot_data["config"]["db"],
                                  "SELECT nomi, summa FROM transactions ORDER BY id")
    data = await asyncio.to_thread(make_excel, rows)
    await update.effective_message.reply_document(document=data,
                                                 filename="kirim_chiqim.xlsx")


async def handle(update, context):
    if not authorized(update, context):
        return
    try:
        name, amount = parse_operation(update.effective_message.text)
    except ValueError as exc:
        await update.effective_message.reply_text(str(exc))
        return
    # Do not retry INSERT: a lost COMMIT response may still mean data was saved.
    await asyncio.to_thread(query, context.bot_data["config"]["db"],
                           "INSERT INTO transactions (nomi, summa) VALUES (%s, %s)",
                           (name, amount), write=True)
    await update.effective_message.reply_text(f"✅ Saqlandi\n{name}: {amount:,} so'm")


async def on_error(update, context):
    exc = context.error
    code = exc.args[0] if isinstance(exc, pymysql.MySQLError) and exc.args else None
    # Never log message contents, credentials, or Telegram token URLs.
    LOG.error("Handler failed: type=%s db_code=%s", type(exc).__name__, code)
    if update and authorized(update, context) and update.effective_message:
        try:
            await update.effective_message.reply_text(
                "❌ Amal yakunlanmadi. /excel orqali yozuvni tekshiring; "
                "takror yuborishdan oldin saqlanganini aniqlang.")
        except Exception as send_exc:
            LOG.error("Error notification failed: %s", type(send_exc).__name__)


def main():
    logging.basicConfig(level=logging.WARNING, format="%(levelname)s %(name)s: %(message)s")
    LOG.setLevel(logging.INFO)
    try:
        config = load_config()
    except ValueError as exc:
        LOG.error("Configuration error: %s", exc)
        raise SystemExit(1) from None
    try:
        app = ApplicationBuilder().token(config["token"]).post_init(initialize).build()
        app.bot_data["config"] = config
        app.add_handler(CommandHandler("start", start))
        app.add_handler(CommandHandler(["balans", "hisobot"], report))
        app.add_handler(CommandHandler(["excel", "exel"], excel))
        app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle))
        app.add_error_handler(on_error)
        app.run_polling(drop_pending_updates=False)
    except Exception as exc:
        code = exc.args[0] if isinstance(exc, pymysql.MySQLError) and exc.args else None
        LOG.error("Startup stopped: type=%s db_code=%s", type(exc).__name__, code)
        raise SystemExit(1) from None


if __name__ == "__main__":
    main()
